"""Consistency checks and derived tables for the multi-cloud governance SLR (v1, ETECOM 2026).

Usage (from repository root):
    pip install -r 08_analysis/requirements.txt
    python 08_analysis/scripts/check_consistency.py

Reads 01_search .. 07_prisma, never modifies them, and writes derived tables to
09_tables_figures/generated/. Exits non-zero if a hard check fails.
"""
import csv, re, sys
from collections import Counter
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "09_tables_figures" / "generated"
DIMS = ["D1", "D2", "D3", "D4", "D5", "D6"]
fails = []

def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        fails.append(msg)

def rows(path, sheet=None):
    wb = openpyxl.load_workbook(ROOT / path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    return [dict(zip(hdr, r)) for r in it if any(c is not None for c in r)]

studies = list(csv.DictReader(open(ROOT / "06_included_studies/included_studies.csv", encoding="utf-8")))
sids = {s["study_id"] for s in studies}
phase1 = [s for s in studies if s["phase"].startswith("Phase 1")]
codes = {r["code_id"] for r in rows("03_coding/codebook.xlsx", "Codebook")}
ev = rows("03_coding/coding_evidence.xlsx")
screen = rows("02_screening/screening.xlsx", "FullText_Screening")
prisma = {r["stage"]: r["count"] for r in rows("07_prisma/prisma_counts.xlsx")}
search = rows("01_search/search_history.xlsx")

# --- referential integrity ---
check(len(studies) == 14 and len(phase1) == 10, "corpus = 10 academic + 4 normative sources")
check(all(r["study_id"] in sids for r in ev), "every coding row references a known study_id")
check(all(r["code_id"] in codes for r in ev), "every coding row references a code_id in the codebook")
cells = Counter((r["study_id"], r["dimension"]) for r in ev)
check(len(cells) == 84 and set(cells.values()) == {1}, "exactly one coded cell per study x dimension (14 x 6)")
inc = {r["study_id_if_included"] for r in screen if r["ft_decision"] == "Include"}
check(inc == {s["study_id"] for s in phase1}, "full-text Include decisions = Phase 1 included studies")
bibkeys = set(re.findall(r"@\w+\{(\w+),", (ROOT / "06_included_studies/included_studies.bib").read_text(encoding="utf-8")))
check(bibkeys == {s["cite_key"] for s in studies}, "included_studies.bib keys = included_studies.csv cite_keys")

# --- PRISMA arithmetic ---
db = [r for r in search if r["search_id"].startswith("SRCH") and r["database"] != "Citation Tracking"]
check(sum(r["hits_returned"] for r in db) == prisma["Records identified from databases"], "database hits sum to PRISMA identification count")
ident = prisma["Records identified from databases"] + prisma["Records identified via other methods"]
check(ident - prisma["Duplicate records removed"] == prisma["Records screened (title/abstract)"], "identified - duplicates = screened")
check(prisma["Records screened (title/abstract)"] - prisma["Records excluded (title/abstract)"] == prisma["Reports assessed for eligibility"], "screened - excluded = assessed for eligibility")
check(sum(1 for r in screen if r["ft_decision"] == "Exclude") == prisma["Reports excluded (full text)"], "full-text exclusions match PRISMA count")
check(prisma["Reports assessed for eligibility"] - prisma["Reports excluded (full text)"] == prisma["Studies included (Phase 1)"] == len(phase1), "assessed - excluded = included")

# --- informational: archived EBSCO export vs. date filter ---
wb = openpyxl.load_workbook(ROOT / "01_search/exports/search_results_by_database_2026-04.xlsx", read_only=True)
eb = [r for r in list(wb["EBSCO"].iter_rows(values_only=True))[1:] if r and r[4]]
yrs = Counter(str(r[6])[:4] for r in eb)
print(f"INFO archived EBSCO records: {len(eb)}; outside 2020-2025: {sum(v for k, v in yrs.items() if not '2020' <= k <= '2025')} ({dict(sorted(yrs.items()))})")

# --- derived tables ---
OUT.mkdir(parents=True, exist_ok=True)
group = {s["study_id"]: ("academic" if s["phase"].startswith("Phase 1") else "normative") for s in studies}
freq = Counter()
for r in ev:
    freq[(r["dimension"], group[r["study_id"]], r["coverage_v1_workbook"].rstrip("*"))] += 1
with open(OUT / "coverage_frequency_v1.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["dimension", "group", "Yes", "Partial", "No"])
    for d in DIMS:
        for g in ("academic", "normative"):
            w.writerow([d, g] + [freq[(d, g, v)] for v in ("Yes", "Partial", "No")])
pub = {(r["study_id"], r["dimension"]): r["coverage_published_etecom2026"] for r in ev if r["coverage_published_etecom2026"]}
with open(OUT / "tableIII_published_etecom2026.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["study_id", "label"] + DIMS)
    for s in studies:
        if group[s["study_id"]] == "normative":
            w.writerow([s["study_id"], s["label"]] + [pub[(s["study_id"], d)] for d in DIMS])
diff = [(k, r["coverage_v1_workbook"], r["coverage_published_etecom2026"]) for r in ev
        for k in [r["evidence_id"]] if r["coverage_published_etecom2026"] and r["coverage_v1_workbook"].rstrip("*") != r["coverage_published_etecom2026"]]
print(f"INFO published Table III differs from v1 coding in {len(diff)} cells: {diff}")
print(f"\n{len(fails)} check(s) failed" if fails else "\nall checks passed")
sys.exit(1 if fails else 0)
